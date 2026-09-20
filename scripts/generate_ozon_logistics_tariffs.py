"""Generate the compact Moscow-origin Ozon logistics tariff bundle.

The source workbook is downloaded from Ozon Seller Education.  Keeping the
conversion deterministic lets us audit a displayed rate against the official
row without shipping a spreadsheet parser in the desktop application.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "desktop-next/src-tauri/resources/ozon-tariffs/logistika-fbo-fbs-28082026.xlsx"
OUTPUT = ROOT / "desktop-next/src/data/ozon-logistics-tariffs-2026-08-28.json"
OFFICIAL_URL = (
    "https://cdn.ozone.ru/s3/ozon-disk-api/Seller-edu/files/"
    "commissions-tariffs/commissions/logistika-fbo-fbs-28082026_1784031962.xlsx"
)
MOSCOW = "Москва, МО и Дальние регионы"


def band_max(value: str) -> float:
    numbers = re.findall(r"\d+(?:[,.]\d+)?", value)
    if not numbers:
        raise ValueError(f"Cannot parse volume band: {value!r}")
    return float(numbers[-1].replace(",", "."))


def main() -> None:
    book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    logistics = book.worksheets[0]
    defaults = book.worksheets[1]

    routes: dict[str, list[dict[str, float]]] = {}
    for row in logistics.iter_rows(min_row=4, values_only=True):
        volume, origin, destination, under_300, over_300 = row[1:6]
        if origin != MOSCOW or not destination or under_300 is None or over_300 is None:
            continue
        routes.setdefault(str(destination), []).append(
            {
                "maxLiters": band_max(str(volume)),
                "under300": float(under_300),
                "over300": float(over_300),
            }
        )

    fallback: list[dict[str, float]] = []
    for row in defaults.iter_rows(min_row=4, values_only=True):
        volume, under_300, over_300 = row[1:4]
        if not volume or under_300 is None or over_300 is None:
            continue
        fallback.append(
            {
                "maxLiters": band_max(str(volume)),
                "under300": float(under_300),
                "over300": float(over_300),
            }
        )

    payload = {
        "effectiveFrom": "2026-08-28",
        "sourceUrl": OFFICIAL_URL,
        "sourceFile": SOURCE.name,
        "origin": MOSCOW,
        "routes": routes,
        "fallback": fallback,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"wrote {OUTPUT} ({len(routes)} destinations, {sum(map(len, routes.values()))} rows)")


if __name__ == "__main__":
    main()
